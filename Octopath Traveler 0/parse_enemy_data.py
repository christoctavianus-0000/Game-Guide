"""Export the workbook's Bestiary sheet for the web app."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

WORKBOOK = Path("Octopath traveler 0 database .xlsx")
OUTPUT_JSON = Path("enemyData.json")
OUTPUT_JS = Path("enemyData.js")

WEAPON_FIELDS = [
    "Sword",
    "Spear",
    "Dagger",
    "Axe",
    "Bow",
    "Staff",
    "Tome",
    "Fan",
    "Fire",
    "Ice",
    "Thunder",
    "Wind",
    "Light",
    "Dark",
]

SKILL_NAME_FIELDS = [f"Skill Name {n}" for n in range(1, 17)]
SKILL_TYPE_FIELDS = [f"Skill Type {n}" for n in range(1, 17)]


def clean_text(value):
    if value is None:
        return ""
    text = str(value).replace("\n", " ")
    return " ".join(text.split())


def to_int(value):
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    if not text:
        return 0
    if text.lower() in {"x", "—", "-"}:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def normalize_weaknesses(row):
    weaknesses = []
    for field in WEAPON_FIELDS:
        value = row.get(field)
        if value is None:
            continue
        text = clean_text(value)
        if not text or text.lower() in {"x", "—", "-"}:
            continue
        weaknesses.append(field)
    return weaknesses


def clean_id(value, fallback_name: str, seen_ids: dict[str, int]) -> str:
    raw = clean_text(value)
    if raw:
        raw = re.sub(r"\.0$", "", raw)
    if not raw or raw.lower() == "xxx":
        raw = re.sub(r"[^a-z0-9]+", "-", fallback_name.lower()).strip("-") or "enemy"
    count = seen_ids.get(raw, 0) + 1
    seen_ids[raw] = count
    return raw if count == 1 else f"{raw}-{count}"


def load_arena_rewards(workbook) -> tuple[dict[str, str], dict[str, str]]:
    """Read Overworld Capture Rewards and Arena Win Rewards from the Monster Arena sheet."""
    capture_rewards: dict[str, str] = {}
    arena_win_rewards: dict[str, str] = {}
    if "Monster Arena" not in workbook.sheetnames:
        return capture_rewards, arena_win_rewards

    ws = workbook["Monster Arena"]
    mode = ""
    for r in ws.iter_rows(values_only=True):
        if not r:
            continue
        col4 = clean_text(r[4]) if len(r) > 4 else ""
        if col4 == "Overworld (Capture)":
            mode = "capture"
            continue
        if col4 == "Arena":
            mode = "arena"
            continue
        if col4 == "StrengthIndex":
            continue

        reward = clean_text(r[12]) if len(r) > 12 else ""
        qty = to_int(r[13]) if len(r) > 13 else 0
        if not reward:
            continue
        reward_str = f"{reward} x{qty}" if qty > 1 else reward

        for col_idx in (5, 6, 7):
            if len(r) > col_idx and r[col_idx]:
                ename = clean_text(r[col_idx]).lower()
                if ename:
                    if mode == "capture":
                        capture_rewards[ename] = reward_str
                    elif mode == "arena":
                        arena_win_rewards[ename] = reward_str

    return capture_rewards, arena_win_rewards


def parse_bestiary(ws, capture_rewards: dict[str, str] | None = None, arena_win_rewards: dict[str, str] | None = None):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = rows[0]
    records = []
    seen_ids: dict[str, int] = {}
    current_section = "Regular"

    for raw in rows[1:]:
        row = {header[idx]: raw[idx] for idx in range(min(len(header), len(raw)))}
        name = clean_text(row.get("Name"))
        area = clean_text(row.get("Area"))
        sub_area = clean_text(row.get("Sub Area"))

        # Track authoritative worksheet section headers
        if not name:
            if area == "Bosses":
                current_section = "Boss"
            elif area in ("Rare", "Sidequest", "Arena", "Elite"):
                current_section = area
            elif sub_area == "XXX" and current_section == "Sidequest" and area.startswith("Bestower"):
                # Rows 388-392 are Sunken Ship sidequest enemies, followed by Row 393 Bestower of All 4
                current_section = "Regular"
            continue

        if not area:
            continue

        weaknesses = normalize_weaknesses(row)
        skills = []
        for name_field, type_field in zip(SKILL_NAME_FIELDS, SKILL_TYPE_FIELDS):
            skill_name = clean_text(row.get(name_field))
            type_name = clean_text(row.get(type_field))
            if skill_name:
                skills.append({"name": skill_name, "type": type_name or "Unknown"})

        enemy_id = clean_id(row.get("ID"), name, seen_ids)
        record = {
            "id": enemy_id,
            "name": name,
            "area": area,
            "subArea": sub_area,
            "type": current_section,
            "shields": to_int(row.get("Shields")),
            "weaknesses": weaknesses,
            "weakness": weaknesses[0] if weaknesses else "None",
            "level": to_int(row.get("Level")),
            "hp": to_int(row.get("HP")),
            "patk": to_int(row.get("P. Atk.")),
            "pdef": to_int(row.get("P. Def.")),
            "matk": to_int(row.get("M. Atk.")),
            "mdef": to_int(row.get("M. Def.")),
            "speed": to_int(row.get("SPD")),
            "crit": to_int(row.get("CRT")),
            "critDef": to_int(row.get("CRTDef")),
            "accuracy": to_int(row.get("ACC")),
            "evasion": to_int(row.get("EVA")),
            "exp": to_int(row.get("EXP")),
            "jp": to_int(row.get("JP")),
            "leaves": to_int(row.get("L")),
            "steal": clean_text(row.get("Steal")),
            "drop": clean_text(row.get("Drop")),
            "immunities": clean_text(row.get("Status Immunities")),
            "skills": skills,
        }

        if current_section == "Arena":
            key_lower = name.lower()
            if area == "Arena":
                win_rew = (arena_win_rewards or {}).get(key_lower, "")
                if not win_rew and key_lower == "immortal flame":
                    win_rew = (arena_win_rewards or {}).get("sultravi", "")
                if win_rew:
                    record["arenaReward"] = f"Arena Win: {win_rew}"
            else:
                cap_rew = (capture_rewards or {}).get(key_lower, "")
                if cap_rew:
                    record["arenaReward"] = f"Capture: {cap_rew}"

        records.append(record)

    return records


_REGIONS = ("Frostlands", "Flatlands", "Woodlands", "Coastlands",
            "Cliftlands", "Riverlands", "Sunlands", "Highlands")


def correct_boss_areas(enemies: list[dict]) -> int:
    """Fix boss areas using subArea→region mapping from regular enemies.

    The workbook stores story-chapter names (e.g. "Master of Wealth Ch. 3")
    as the area for story bosses, but the codex should show the geographic
    region (e.g. "Woodlands").  This cross-references each boss's subArea
    against regular-enemy subAreas to recover the correct region.
    Returns the number of bosses corrected.
    """
    # 1. Build subArea → region from regular (non-boss) enemies only.
    sub_to_region: dict[str, str] = {}
    for e in enemies:
        if e["type"] in ("Boss", "Rare", "Arena", "Sidequest"):
            continue
        sa = e.get("subArea", "").strip()
        area = e.get("area", "").strip()
        if sa and area in _REGIONS:
            key = sa.lower()
            if key not in sub_to_region:
                sub_to_region[key] = area

    # 2. Correct boss areas.
    fixed = 0
    for e in enemies:
        if e["type"] != "Boss":
            continue
        current_area = e.get("area", "")
        # Only fix bosses whose area is a story-chapter / meta label
        if current_area in _REGIONS:
            continue
        sa = e.get("subArea", "").strip()
        if not sa:
            continue
        region = sub_to_region.get(sa.lower())
        if region and region != current_area:
            e["area"] = region
            fixed += 1
    return fixed


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    capture_rewards, arena_win_rewards = load_arena_rewards(workbook)
    ws = workbook["Bestiary"]
    enemies = parse_bestiary(ws, capture_rewards, arena_win_rewards)

    fixed = correct_boss_areas(enemies)
    if fixed:
        print(f"Corrected {fixed} boss areas via subArea cross-reference.")

    OUTPUT_JSON.write_text(json.dumps(enemies, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    OUTPUT_JS.write_text(f"const enemyData = {json.dumps(enemies, ensure_ascii=False, indent=2)};\n", encoding="utf-8")

    print(f"Exported {len(enemies)} enemies from {WORKBOOK}.")


if __name__ == "__main__":
    main()

