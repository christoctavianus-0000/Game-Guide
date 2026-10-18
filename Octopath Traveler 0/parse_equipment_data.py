"""Export the workbook's player-facing Weapon and Armor sheets for the web app."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path("Octopath traveler 0 database .xlsx")
OUTPUT_JSON = Path("equipmentData.json")
OUTPUT_JS = Path("equipmentData.js")

STAT_FIELDS = ("HP", "SP", "PATK", "PDEF", "MATK", "MDEF", "SPD", "CRT", "ACC", "EVA")

WEAPON_ICONS = {
    "Sword": "assets/icons/icon_sword.png",
    "Polearm": "assets/icons/icon_polearm.png",
    "Dagger": "assets/icons/icon_dagger.png",
    "Axe": "assets/icons/icon_axe.png",
    "Bow": "assets/icons/icon_bow.png",
    "Staff": "assets/icons/icon_staff.png",
    "Tome": "assets/icons/icon_tome.png",
    "Fan": "assets/icons/icon_fan.png",
}


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def clean_number(value: object) -> int | float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def clean_cost(value: object) -> tuple[str, int | None]:
    """Keep multi-source prices readable while supplying a stable sort value."""
    text = clean_text(value)
    if not text:
        return "", None
    amounts = [int(amount.replace(",", "")) for amount in re.findall(r"\d[\d,]*", text)]
    return text, min(amounts) if amounts else None


def build_source_backfill(wb) -> dict[str, str]:
    """Cross-reference Treasure, Sidequests, Townspeople, and Bestiary sheets for items with blank SOURCE."""
    backfill: dict[str, list[str]] = {}

    def add_src(item_name: str, src: str) -> None:
        key = item_name.strip().lower()
        if not key:
            return
        lst = backfill.setdefault(key, [])
        if src not in lst:
            lst.append(src)

    if "Treasure" in wb.sheetnames:
        for r in wb["Treasure"].iter_rows(min_row=2, values_only=True):
            if not r or not r[0]:
                continue
            what = clean_text(r[0])
            area = clean_text(r[1]) if len(r) > 1 else ""
            sub = clean_text(r[2]) if len(r) > 2 else ""
            notes = clean_text(r[3]) if len(r) > 3 else ""
            if what and (area or sub):
                loc = f"{sub} ({area})" if sub and area else (sub or area)
                detail = f"[Treasure] {loc}" + (f" - {notes}" if notes else "")
                add_src(what, detail)

    if "Sidequests" in wb.sheetnames:
        for r in wb["Sidequests"].iter_rows(min_row=3, values_only=True):
            if not r or len(r) < 7 or not r[1] or not r[6]:
                continue
            qname = clean_text(r[1])
            qloc = clean_text(r[2])
            rewards = [line.strip() for line in str(r[6]).splitlines() if line.strip()]
            for rew in rewards:
                add_src(rew, f"[Sidequest] {qname} ({qloc})")

    if "Townspeople" in wb.sheetnames:
        for r in wb["Townspeople"].iter_rows(min_row=4, values_only=True):
            if not r or not r[0]:
                continue
            npc = clean_text(r[0])
            loc = clean_text(r[1]) if len(r) > 1 else ""
            if not loc:
                continue
            inquire_item = clean_text(r[4]) if len(r) > 4 else ""
            if inquire_item and inquire_item.lower() not in ("none", "n/a"):
                add_src(inquire_item, f"[NPC] {loc} - {npc} (Inquire)")
            for ci in range(13, min(18, len(r))):
                itm = clean_text(r[ci])
                if itm and itm.lower() not in ("none", "n/a"):
                    add_src(itm, f"[NPC] {loc} - {npc}")

    if "Bestiary" in wb.sheetnames:
        for r in wb["Bestiary"].iter_rows(min_row=2, values_only=True):
            if not r or len(r) < 35 or not r[3] or not r[0]:
                continue
            ename = clean_text(r[3])
            sub = clean_text(r[1]) or clean_text(r[0])
            steal = clean_text(r[33])
            if steal and steal.lower() not in ("none", "n/a", "-"):
                add_src(steal, f"[Steal] {ename} ({sub})")

    return {k: ", ".join(v) for k, v in backfill.items()}


def read_sheet(wb, sheet_name: str, name_column: str, kind: str, source_backfill: dict[str, str]) -> list[dict]:
    # Rows 1-3 are workbook notes/spacing; row 4 contains the actual headers.
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    if not rows:
        return []
    headers = [clean_text(c) for c in rows[0]]
    records: list[dict] = []

    for row_number, raw in enumerate(rows[1:], start=1):
        row = {headers[i]: raw[i] for i in range(min(len(headers), len(raw))) if headers[i]}
        name = clean_text(row.get(name_column))
        item_type = clean_text(row.get("TYPE"))
        if not name or not item_type:
            continue

        stats = {field.lower(): clean_number(row.get(field)) or 0 for field in STAT_FIELDS}
        cost_raw = row.get("COST")
        if isinstance(cost_raw, float) and cost_raw.is_integer():
            cost_raw = int(cost_raw)
        cost, cost_value = clean_cost(cost_raw)
        source = clean_text(row.get("SOURCE"))
        if not source:
            source = source_backfill.get(name.lower(), "")

        records.append(
            {
                "id": f"{kind.lower()}-{row_number}",
                "recordType": "equipment",
                "kind": kind,
                "name": name,
                "type": item_type,
                "icon": WEAPON_ICONS.get(item_type, "") if kind == "Weapon" else "",
                "stats": stats,
                "effect": clean_text(row.get("EFFECT")),
                "source": source,
                "cost": cost,
                "costValue": cost_value,
            }
        )

    return records


def main() -> None:
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    source_backfill = build_source_backfill(wb)
    weapons = read_sheet(wb, "Weapon", "WEAPON", "Weapon", source_backfill)
    armor = read_sheet(wb, "Armor", "Equipment", "Armor", source_backfill)
    wb.close()
    equipment = weapons + armor

    serialized = json.dumps(equipment, ensure_ascii=False, indent=2)
    OUTPUT_JSON.write_text(serialized + "\n", encoding="utf-8")
    OUTPUT_JS.write_text(f"const equipmentData = {serialized};\n", encoding="utf-8")

    print(f"Exported {len(weapons)} weapons and {len(armor)} armor pieces.")


if __name__ == "__main__":
    main()

