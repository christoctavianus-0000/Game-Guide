"""Parse mastery skill sources from the workbook using cell background colors.

Color mapping (from row-1 headers):
  Yellow (FFFFF2CC) = Chest
  Red    (FFF4CCCC) = Boss
  Blue   (FFCFE2F3) = NPC
  Green  (FFD9EAD3) = Trade Goods
  No fill / other   = Other
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

WORKBOOK = Path("Octopath traveler 0 database .xlsx")
OUTPUT_JS = Path("masterySourceData.js")

# Source columns I-L (1-indexed: 9-12)
SOURCE_COLS = [9, 10, 11, 12]

# Color → category mapping (openpyxl fgColor.rgb strings, leading FF stripped by some themes)
COLOR_MAP: dict[str, str] = {
    "FFFFF2CC": "Chest",
    "FFF4CCCC": "Boss",
    "FFCFE2F3": "NPC",
    "FFD9EAD3": "Trade Goods",
}


def rgb_category(rgb: str | None, fill_type: str | None) -> str:
    """Map a cell's fgColor.rgb to a category name."""
    if not rgb or fill_type in (None, "none"):
        return "Other"
    return COLOR_MAP.get(rgb, "Other")


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ")
    return " ".join(text.split())


def parse_mastery_sheet(ws) -> dict[str, list[dict]]:
    """Read Source 1-4 columns, categorize by cell background color."""
    sources: dict[str, list[dict]] = {}

    for row_idx in range(5, ws.max_row + 1):
        name = clean_text(ws.cell(row=row_idx, column=7).value)  # col G = Mastery Name
        if not name:
            continue

        seen: set[str] = set()
        entries: list[dict] = []

        for col_idx in SOURCE_COLS:
            cell = ws.cell(row=row_idx, column=col_idx)
            text = clean_text(cell.value)
            if not text or text in (".", "r", " "):
                continue

            fill = cell.fill
            fg = fill.fgColor
            rgb = fg.rgb if fg else None
            category = rgb_category(rgb, fill.patternType)

            # Deduplicate within the same skill
            key = (text.lower(), category)
            if key in seen:
                continue
            seen.add(key)
            entries.append({"text": text, "category": category})

        if entries:
            sources[name] = entries

    return sources


def main():
    wb = load_workbook(WORKBOOK)
    ws = wb.worksheets[16]  # sheet 17 = index 16 = "Mastery Skills"
    sources = parse_mastery_sheet(ws)

    # Build JS output
    lines = ["const masterySourceIndex = {"]
    items = sorted(sources.items())
    for i, (name, entries) in enumerate(items):
        json_entries = json.dumps(entries, ensure_ascii=False)
        comma = "," if i < len(items) - 1 else ""
        lines.append(f"  {json.dumps(name)}:{json_entries}{comma}")
    lines.append("};")
    lines.append("")

    OUTPUT_JS.write_text("\r\n".join(lines), encoding="utf-8")
    print(f"Wrote {len(sources)} mastery skills to {OUTPUT_JS}")

    # Print category stats
    from collections import Counter
    cats: Counter[str] = Counter()
    for entries in sources.values():
        for e in entries:
            cats[e["category"]] += 1
    print("Category distribution:")
    for cat, count in cats.most_common():
        print(f"  {cat}: {count}")


if __name__ == "__main__":
    main()
