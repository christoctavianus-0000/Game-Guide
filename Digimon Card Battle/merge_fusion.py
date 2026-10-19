"""Merge Fusion Material/Result values into card JSON files."""
import json
import openpyxl
import os

BASE = os.path.dirname(os.path.abspath(__file__))

# Load fusion values
wb = openpyxl.load_workbook(os.path.join(BASE, 'Digimon Card Battle - Fusion.xlsx'))
ws = wb['Fusion value']
fusion = {}
fusion_rows = []
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    color, name, material, result = row
    if name:
        name = name.strip()
        fusion[name.lower()] = (material, result)
        fusion_rows.append({
            'Color': color,
            'Name': name,
            'Material': material,
            'Result': result
        })

with open(os.path.join(BASE, 'fusion_values.json'), 'w', encoding='utf-8') as f:
    json.dump(fusion_rows, f, indent=2, ensure_ascii=False)

# Files to update
files = [
    'fireCards.json', 'iceCards.json', 'natureCards.json', 'darknessCards.json',
    'rareCards.json', 'partnerCards.json', 'optionCards.json'
]

# Map canonical in-game card names to spelling variants in the GameFAQs guide / workbook
NAME_ALIASES = {
    'phoenixmon': 'pheonixmon',
    'gigadramon': 'gigidramon',
    'monochromon': 'monocromon',
    'brachiomon': 'braciomon',
    'silphymon': 'sylphymon',
    'ogremon': 'orgemon',
    "dark lord's cape": 'dark lord cape',
    'ice altar': 'ice alter',
    'nature altar': 'nature alter',
    'darkness altar': 'darkness alter',
    'rare altar': 'rare alter',
    'mega def. disk ○': 'mega def. disk circle',
    'mega def. disk △': 'mega def. disk triangle',
    'mega def. disk ×': 'mega def. disk x',
    "suka's curse": 'sukas curse',
    "cherrymon's mist": 'cherrymons mist',
    'digi-sapphire': 'digi-saphire',
    'speed digivolve': 'speed-digivolve',
}

# Map disk cards by card number to avoid character encoding mismatches
CARD_NO_ALIASES = {
    '222': 'mega def. disk circle',
    '223': 'mega def. disk triangle',
    '224': 'mega def. disk x',
    '267': 'attack disk circle',
    '268': 'attack disk triangle',
    '269': 'attack disk x',
    '270': 'defense disk circle',
    '271': 'defense disk triangle',
    '272': 'defense disk x',
}

for fname in files:
    path = os.path.join(BASE, fname)
    with open(path, 'r', encoding='utf-8') as f:
        cards = json.load(f)
    updated = 0
    for card in cards:
        c_no = card.get('no')
        raw_key = card.get('name', '').strip().lower()
        if c_no in CARD_NO_ALIASES:
            key = CARD_NO_ALIASES[c_no]
        else:
            key = NAME_ALIASES.get(raw_key, raw_key)
        if key in fusion:
            mat, res = fusion[key]
            card['materialValue'] = mat
            card['resultValue'] = res
            updated += 1
        else:
            card['materialValue'] = '?'
            card['resultValue'] = '?'
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(cards, f, indent=2)
    print(f'{fname}: {updated}/{len(cards)} matched')