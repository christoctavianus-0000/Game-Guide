from openpyxl import load_workbook
from pathlib import Path

wb = load_workbook(r'd:\My AI Project\Game Guide\Octopath traveler 0 database .xlsx', read_only=False, data_only=True)
ws = wb['Character']
rows = list(ws.iter_rows(values_only=True))

# print all character entries with battle skills, and Alfyn skills exactly as workbook has them
for row in rows:
    if len(row) > 1 and row[1] is not None:
        label = str(row[1]).replace('\n',' ').replace('  ',' ').strip()
        if label.startswith('Alfyn'):
            print('ALFYN ROW LABEL:', label)
            for i, cell in enumerate(row):
                if i < 12:
                    print(i, repr(cell))
            break

print('--- battle rows ---')
for row in rows:
    if len(row) >= 6:
        type_val = row[2] if len(row) > 2 else None
        name_val = row[4] if len(row) > 4 else None
        if type_val == 'Battle' and name_val:
            print('BATTLE ROW:', repr(type_val), repr(name_val), 'ROW', rows.index(row))

print('--- character entries ---')
seen = []
for row in rows:
    if len(row) > 1 and row[1] is not None:
        label = str(row[1]).replace('\n',' ').replace('  ',' ').strip()
        if label and ' - ' in label and '(' in label and ')' in label:
            seen.append(label)
print(seen[:10])
print('count', len(seen))
