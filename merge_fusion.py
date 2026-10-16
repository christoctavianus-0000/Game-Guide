"""Merge Fusion Material/Result values into card JSON files."""
import json
import openpyxl
import os

BASE = os.path.dirname(os.path.abspath(__file__))

# Load fusion values
wb = openpyxl.load_workbook(os.path.join(BASE, 'Digimon Card Battle', 'Digimon Card Battle - Fusion.xlsx'))
ws = wb['Fusion value']
fusion = {}
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    color, name, material, result = row
    if name:
        fusion[name.strip().lower()] = (material, result)

# Files to update
files = [
    'fireCards.json', 'iceCards.json', 'natureCards.json', 'darknessCards.json',
    'rareCards.json', 'partnerCards.json', 'optionCards.json'
]

for fname in files:
    path = os.path.join(BASE, 'Digimon Card Battle', fname)
    with open(path, 'r') as f:
        cards = json.load(f)
    updated = 0
    for card in cards:
        key = card.get('name', '').strip().lower()
        if key in fusion:
            mat, res = fusion[key]
            card['materialValue'] = mat
            card['resultValue'] = res
            updated += 1
        else:
            card['materialValue'] = '?'
            card['resultValue'] = '?'
    with open(path, 'w') as f:
        json.dump(cards, f, indent=2)
    print(f'{fname}: {updated}/{len(cards)} matched')