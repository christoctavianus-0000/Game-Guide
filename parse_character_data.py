"""Export the workbook's Character and Protagonist Job sheets for the web app."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

WORKBOOK = Path("Octopath traveler 0 database .xlsx")
OUTPUT_JSON = Path("characterData.json")
OUTPUT_JS = Path("characterData.js")


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ")
    return " ".join(text.split())


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return slug or "source"


def split_character_label(label: str) -> tuple[str, str, str]:
    text = label.strip()
    if not text:
        return "", "", ""

    if "(" in text and ")" in text:
        head = text[: text.rfind("(")].strip()
        weapons = text[text.rfind("(") + 1 : text.rfind(")")].strip()
    else:
        head = text
        weapons = ""

    if re.search(r"\s*[-–]\s*", head):
        name, job = re.split(r"\s*[-–]\s*", head, maxsplit=1)
        return name.strip(), job.strip(), weapons

    return head.strip(), "", weapons


def parse_character_sheet(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    characters: list[dict] = []
    current: dict | None = None
    pending_unlock = ""
    pending_amity_parts: list[str] = []
    current_skill_type: str | None = None

    for row in rows:
        col_b = row[1] if len(row) > 1 else None
        col_c = row[2] if len(row) > 2 else None
        col_d = row[3] if len(row) > 3 else None
        col_e = row[4] if len(row) > 4 else None
        col_f = row[5] if len(row) > 5 else None
        col_g = row[6] if len(row) > 6 else None
        col_h = row[7] if len(row) > 7 else None
        col_i = row[8] if len(row) > 8 else None

        raw_b = str(col_b) if col_b is not None else ""
        label = clean_text(raw_b)

        if label.startswith("Unlock Requirements:"):
            pending_unlock = label.replace("Unlock Requirements:", "").strip()
            pending_amity_parts = []
            current_skill_type = None
            continue

        if label.lower().startswith(("character - job", "job", "notes", "the jobs below")):
            current_skill_type = None
            continue

        is_character_entry = (
            label
            and "(" in label
            and ")" in label
            and re.search(r"\s*[-–]\s*", label)
            and not label.lower().startswith("unlock requirements:")
        )

        if is_character_entry:
            if current is not None:
                # Convert amity list to string before appending
                if "amity" in current and isinstance(current["amity"], list):
                    current["amity"] = "; ".join(current["amity"])
                elif "amity" not in current:
                    current["amity"] = ""
                characters.append(current)
            name, job_name, weapons = split_character_label(label)
            current = {
                "id": slugify(name),
                "name": name,
                "title": job_name,
                "job": job_name,
                "weapons": weapons,
                "unlock": pending_unlock,
                "skills": [],
            }
            pending_unlock = ""
            current_skill_type = None

        if current is None:
            continue

        skill_type = clean_text(col_c)
        if skill_type in {"Battle", "Support", "Ultimate"}:
            current_skill_type = skill_type

        skill_name = clean_text(col_e)
        if skill_name and (skill_type in {"Battle", "Support", "Ultimate"} or current_skill_type is not None):
            final_skill_type = skill_type if skill_type in {"Battle", "Support", "Ultimate"} else current_skill_type
            current["skills"].append(
                {
                    "name": skill_name,
                    "type": final_skill_type,
                    "cost": clean_text(col_f),
                    "description": clean_text(col_g),
                    "unmasterable": bool(col_h),
                    "other": clean_text(col_i),
                }
            )

        # Collect amity info from column I (index 8) for skill rows
        if current is not None and col_i:
            amity_text = clean_text(col_i)
            if amity_text and amity_text != 'Amity' and amity_text != 'Other':
                if "amity" not in current:
                    current["amity"] = []
                current["amity"].append(amity_text)

        if is_character_entry:
            continue

    if current is not None:
        # Convert amity list to string before appending
        if "amity" in current and isinstance(current["amity"], list):
            current["amity"] = "; ".join(current["amity"])
        elif "amity" not in current:
            current["amity"] = ""
        characters.append(current)

    return characters


def parse_job_sheet(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    jobs: list[dict] = []
    current: dict | None = None
    current_skill_type: str | None = None

    for row in rows:
        col_b = row[1] if len(row) > 1 else None
        col_c = row[2] if len(row) > 2 else None
        col_d = row[3] if len(row) > 3 else None
        col_e = row[4] if len(row) > 4 else None
        col_f = row[5] if len(row) > 5 else None
        col_g = row[6] if len(row) > 6 else None
        col_h = row[7] if len(row) > 7 else None
        col_i = row[8] if len(row) > 8 else None

        raw_b = str(col_b) if col_b is not None else ""
        label = clean_text(raw_b)

        if label.lower().startswith(("notes", "job", "the jobs below")):
            current_skill_type = None
            continue

        is_job_entry = label and "(" in label and ")" in label and not label.lower().startswith(("notes", "job"))

        if is_job_entry:
            if current is not None:
                jobs.append(current)
            name = label.split("(", 1)[0].strip()
            weapons = label[label.find("(") + 1 : label.rfind(")")].strip()
            current = {
                "id": slugify(name),
                "name": name,
                "weapons": weapons,
                "skills": [],
            }
            current_skill_type = None

        if current is None:
            continue

        skill_type = clean_text(col_c)
        if skill_type in {"Battle", "Support", "Ultimate"}:
            current_skill_type = skill_type

        skill_name = clean_text(col_e)
        if skill_name and (skill_type in {"Battle", "Support", "Ultimate"} or current_skill_type is not None):
            final_skill_type = skill_type if skill_type in {"Battle", "Support", "Ultimate"} else current_skill_type
            current["skills"].append(
                {
                    "name": skill_name,
                    "type": final_skill_type,
                    "cost": clean_text(col_f),
                    "description": clean_text(col_g),
                    "unlock": clean_text(col_h),
                    "unmasterable": bool(col_i),
                }
            )

        if is_job_entry:
            continue

    if current is not None:
        jobs.append(current)

    return jobs


def build_protagonist_entry(jobs: list[dict]) -> dict:
    aggregate_skills: list[dict] = []
    for job in jobs:
        for skill in job.get("skills", []):
            entry = dict(skill)
            entry["sourceJob"] = job["name"]
            entry["sourceJobId"] = job["id"]
            aggregate_skills.append(entry)

    weapon_summary = ", ".join(
        job["weapons"] for job in jobs if job.get("weapons")
    ) or "All weapons"
    return {
        "id": "protagonist",
        "name": "Protagonist",
        "title": "Hero",
        "job": "Protagonist",
        "weapons": weapon_summary,
        "unlock": "Acts as the party's main character and can switch between the available protagonist jobs.",
        "skills": aggregate_skills,
    }


def build_source_index(characters: list[dict], jobs: list[dict]) -> dict[str, list[dict]]:
    sources: dict[str, list[dict]] = {}

    for character in characters:
        for skill in character["skills"]:
            name = skill["name"]
            entry = {
                "type": "character",
                "id": character["id"],
                "label": f"{character['name']} - {character['title']}",
                "href": f"character.html#{character['id']}",
                "skillType": skill["type"],
            }
            sources.setdefault(name, []).append(entry)

    for job in jobs:
        for skill in job["skills"]:
            name = skill["name"]
            entry = {
                "type": "job",
                "id": job["id"],
                "label": job["name"],
                "href": f"character.html#job-{job['id']}",
                "skillType": skill["type"],
            }
            sources.setdefault(name, []).append(entry)

    return sources


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=False, data_only=True)
    characters = parse_character_sheet(workbook["Character"])
    jobs = parse_job_sheet(workbook["Protag Job"])
    index = build_source_index(characters, jobs)

    payload = {
        "characters": characters,
        "jobs": jobs,
        "protagonist": build_protagonist_entry(jobs),
        "sourceIndex": index,
    }
    serialized = json.dumps(payload, ensure_ascii=False, indent=2)
    OUTPUT_JSON.write_text(serialized + "\n", encoding="utf-8")
    OUTPUT_JS.write_text(f"const characterData = {serialized};\nconst characterSourceIndex = characterData.sourceIndex;\n", encoding="utf-8")
    print(f"Exported {len(characters)} characters and {len(jobs)} jobs, with {len(index)} skill source keys.")


if __name__ == "__main__":
    main()
