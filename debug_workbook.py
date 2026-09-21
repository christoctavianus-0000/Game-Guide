from openpyxl import load_workbook

path = r'd:\My AI Project\Game Guide\Octopath traveler 0 database .xlsx'
wb = load_workbook(path, read_only=True, data_only=True)
print('SHEETS:', wb.sheetnames)
for name in wb.sheetnames:
    if 'character' in name.lower() or 'job' in name.lower() or 'protag' in name.lower():
        ws = wb[name]
        print(f'\n### {name} rows={ws.max_row} cols={ws.max_column}')
        for row in ws.iter_rows(min_row=1, max_row=min(8, ws.max_row), values_only=True):
            print(row)
