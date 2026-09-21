"""Export the workbook's Bestiary sheet for the web app."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

WORKBOOK = Path("Octopath traveler 0 database .xlsx")
OUTPUT_JSON = Path("enemyData.json")
OUTPUT_JS = Path("enemyData.js")

WEAPON_FIELDS = [
    "Sword",
    "Spear",
    "Dagger",
    "Axe",
    "Bow",
    "Staff",
    "Tome",
    "Fan",
    "Fire",
    "Ice",
    "Thunder",
    "Wind",
    "Light",
    "Dark",
]

SKILL_NAME_FIELDS = [f"Skill Name {n}" for n in range(1, 17)]
SKILL_TYPE_FIELDS = [f"Skill Type {n}" for n in range(1, 17)]


def clean_text(value):
    if value is None:
        return ""
    text = str(value).replace("\n", " ")
    return " ".join(text.split())


def to_int(value):
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    if not text:
        return 0
    if text.lower() in {"x", "—", "-"}:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def normalize_weaknesses(row):
    weaknesses = []
    for field in WEAPON_FIELDS:
        value = row.get(field)
        if value is None:
            continue
        text = clean_text(value)
        if not text or text.lower() in {"x", "—", "-"}:
            continue
        weaknesses.append(field)
    return weaknesses


_SIDEQUEST_AREAS = (
    "an encounter with", "a long-lost love", "alpaca roundup",
    "an assistant like no other", "an indestructible flame",
    "great expectations", "heidne and the", "in search of what remains",
    "one's true worth", "the big 'un", "the blade of the weak",
    "the eight sacred treasures", "the missing brother",
    "treasure event battle", "unnatural tremors",
    "what remains of the black scarlet",
)

_BOSS_AREAS = (
    "master of", "bestower of", "superboss",
    "gate of finis", "finale", "kindlers", "unlocking the true",
    "the heavens",
)

_BOSS_NAMES = (
    "guardian", "dragon", "lord", "demon", "tyrant", "fiend",
    "emperor", "empress", "scholarking", "birdking", "lizardking",
    "ratking", "frogking", "yak", "archmage", "warmaster",
    "runelord", "starseer",
)


def infer_enemy_type(name: str, area: str = "", sub_area: str = "") -> str:
    nl = name.lower()
    al = area.lower()
    sl = sub_area.lower()

    # 1. Rare — Cait & Octopuff variants (high EXP/JP, flee quickly)
    if any(k in nl for k in ("cait", "octopuff")):
        return "Rare"

    # 2. Arena — Monster Arena encounters
    if al == "arena" or "arena" in sl:
        return "Arena"

    # 3. Sidequest — Traveler quests & Treasure Event Battles
    if any(al.startswith(p) for p in _SIDEQUEST_AREAS):
        return "Sidequest"

    # 4. Elite — Enhanced overworld spawns (Menacing, Mighty, Ancient, Greater, etc.)
    if any(k in nl for k in ("menacing", "mighty", "ancient", "greater", "elite")):
        return "Elite"

    # 5. Boss — Story chapters, named boss enemies
    if any(k in al for k in _BOSS_AREAS):
        return "Boss"
    if any(k in nl for k in _BOSS_NAMES):
        return "Boss"
    if any(k in nl for k in ("king", "queen", "general", "boss")):
        return "Boss"
    if "fallen" in nl:
        return "Boss"

    # 6. Regular — standard overworld / dungeon field encounters
    return "Regular"


def parse_bestiary(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = rows[0]
    label_map = {clean_text(cell): idx for idx, cell in enumerate(header)}
    records = []

    for raw in rows[1:]:
        row = {header[idx]: raw[idx] for idx in range(min(len(header), len(raw)))}
        name = clean_text(row.get("Name"))
        area = clean_text(row.get("Area"))
        sub_area = clean_text(row.get("Sub Area"))
        if not name or not area:
            continue

        weaknesses = normalize_weaknesses(row)
        skills = []
        for name_field, type_field in zip(SKILL_NAME_FIELDS, SKILL_TYPE_FIELDS):
            skill_name = clean_text(row.get(name_field))
            type_name = clean_text(row.get(type_field))
            if skill_name:
                skills.append({"name": skill_name, "type": type_name or "Unknown"})

        records.append(
            {
                "id": clean_text(row.get("ID")) or re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-") or "enemy",
                "name": name,
                "area": area,
                "subArea": sub_area,
                "type": infer_enemy_type(name, area, sub_area),
                "shields": to_int(row.get("Shields")),
                "weaknesses": weaknesses,
                "weakness": weaknesses[0] if weaknesses else "None",
                "level": to_int(row.get("Level")),
                "hp": to_int(row.get("HP")),
                "patk": to_int(row.get("P. Atk.")),
                "pdef": to_int(row.get("P. Def.")),
                "matk": to_int(row.get("M. Atk.")),
                "mdef": to_int(row.get("M. Def.")),
                "speed": to_int(row.get("SPD")),
                "crit": to_int(row.get("CRT")),
                "accuracy": to_int(row.get("ACC")),
                "evasion": to_int(row.get("EVA")),
                "exp": to_int(row.get("EXP")),
                "jp": to_int(row.get("JP")),
                "steal": clean_text(row.get("Steal")),
                "drop": clean_text(row.get("Drop")),
                "immunities": clean_text(row.get("Status Immunities")),
                "skills": skills,
            }
        )

    return records


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    ws = workbook["Bestiary"]
    enemies = parse_bestiary(ws)

    OUTPUT_JSON.write_text(json.dumps(enemies, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    OUTPUT_JS.write_text(f"const enemyData = {json.dumps(enemies, ensure_ascii=False, indent=2)};\n", encoding="utf-8")

    print(f"Exported {len(enemies)} enemies from {WORKBOOK}.")


if __name__ == "__main__":
    main()
